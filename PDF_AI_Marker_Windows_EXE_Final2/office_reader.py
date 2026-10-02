# -*- coding: utf-8 -*-
"""
Đọc trực tiếp hồ sơ Word (.docx) và Excel (.xlsx / .xlsm) -> cùng định dạng kết
quả với đường PDF (markdown + bảng + JSON), KHÔNG OCR.

Chỉ dùng python-docx / openpyxl đã có sẵn trong engine, nên vẫn chạy 100% offline.

Khác biệt quan trọng so với PDF:
  * Chữ và số được đọc thẳng từ file nên chính xác tuyệt đối (không có bước đoán).
  * Excel: số được ghi dạng chuẩn dấu chấm thập phân, KHÔNG có dấu phân nghìn
    (1525.81 = 1.525,81). "values" trong JSON là số thực đúng như trong ô, không
    qua bước đoán kiểu VN/US. number_style = "native".
  * Ô gộp: gộp dọc thì lặp giá trị xuống các dòng (mỗi dòng tự đủ nghĩa); gộp ngang
    ở dòng tiêu đề thì lặp sang các cột để tiêu đề nhiều tầng không bị mất nhãn.
  * Ô công thức chưa có giá trị đã tính (file chưa từng mở/lưu bằng Excel) sẽ để
    trống và được cảnh báo, không tự tính lại.
"""
import datetime
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from layout_reconstructor import (
    detect_number_style, fold_key, make_table, parse_number, split_header, table_to_markdown,
)

WORD_EXT = {".docx"}
EXCEL_EXT = {".xlsx", ".xlsm"}
LEGACY_EXT = {".doc": "Word", ".xls": "Excel"}
OFFICE_EXT = WORD_EXT | EXCEL_EXT | set(LEGACY_EXT)


def is_office(path) -> bool:
    return Path(path).suffix.lower() in OFFICE_EXT


def _legacy_error(ext: str) -> ValueError:
    app = LEGACY_EXT[ext]
    return ValueError(
        f"File {ext} là định dạng {app} đời cũ, chưa đọc trực tiếp được. "
        f"Mở bằng {app} → Lưu thành {'.docx' if ext == '.doc' else '.xlsx'} rồi chuyển đổi lại "
        f"(nhiều phần mềm dự toán xuất file .xls thực chất là bảng HTML, mở bằng Excel rồi lưu lại là được).")


def _ws(text: Any) -> str:
    return re.sub(r"\s+", " ", str(text).replace(" ", " ")).strip()


# ─────────────────────────────────────────────────────────────────────────────
# EXCEL
# ─────────────────────────────────────────────────────────────────────────────
def _fmt_native(v: Any) -> str:
    """Chuỗi hiển thị của giá trị ô: số ở dạng chuẩn (không phân nghìn, chấm thập phân)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:
            return str(int(v))
        s = f"{v:.10g}"
        return repr(v) if "e" in s else s
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d") if not (v.hour or v.minute or v.second) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, datetime.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M")
    return _ws(v)


def _is_native_number(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _sheet_matrix(ws_val, ws_formula):
    """(vals, hfill, n_formula_uncalc) — vals: giá trị gốc theo ô (đã lặp ô gộp dọc);
    hfill: bản đã lặp cả ô gộp ngang (chỉ dùng dựng tiêu đề cột)."""
    rows_v = [list(r) for r in ws_val.iter_rows(values_only=True)]
    rows_f = [list(r) for r in ws_formula.iter_rows(values_only=True)]
    nrow = len(rows_v)
    ncol = max((len(r) for r in rows_v), default=0)
    vals = [r + [None] * (ncol - len(r)) for r in rows_v]
    hfill = [list(r) for r in vals]

    uncalc = 0
    for rv, rf in zip(vals, rows_f):
        for a, b in zip(rv, rf):
            if a is None and isinstance(b, str) and b.startswith("="):
                uncalc += 1

    for rng in ws_val.merged_cells.ranges:
        r0, r1, c0, c1 = rng.min_row - 1, rng.max_row - 1, rng.min_col - 1, rng.max_col - 1
        if r0 >= nrow or c0 >= ncol:
            continue
        top = vals[r0][c0]
        for r in range(r0, min(r1, nrow - 1) + 1):
            for c in range(c0, min(c1, ncol - 1) + 1):
                if (r, c) == (r0, c0):
                    continue
                vals[r][c] = top if (c == c0 and r > r0) else None
                hfill[r][c] = top
    return vals, hfill, uncalc


def _blocks(nonempty_rows: List[int]) -> List[List[int]]:
    blocks: List[List[int]] = []
    for r in nonempty_rows:
        if blocks and r == blocks[-1][-1] + 1:
            blocks[-1].append(r)
        else:
            blocks.append([r])
    return blocks


def _excel_block(vals, hfill, rows: List[int]) -> Tuple[List[str], Optional[Dict[str, Any]], List[str], str]:
    """Trả (chữ_trước_bảng, bảng|None, chữ_sau_bảng, tiêu_đề_bảng)."""
    def cells(r):
        return [(c, v) for c, v in enumerate(vals[r]) if v is not None and _fmt_native(v) != ""]

    counts = {r: len(cells(r)) for r in rows}
    multi = [r for r in rows if counts[r] >= 2]

    def as_text(r):
        return _ws(" ".join(_fmt_native(v) for _, v in cells(r)))

    if len(multi) < 2:
        return [as_text(r) for r in rows], None, [], ""

    i0, i1 = rows.index(multi[0]), rows.index(multi[-1])
    lead, body_rows, trail = rows[:i0], rows[i0:i1 + 1], rows[i1 + 1:]
    used = sorted({c for r in body_rows for c, _ in cells(r)})
    c_min, c_max = used[0], used[-1]
    keep = [c for c in range(c_min, c_max + 1) if any(_fmt_native(vals[r][c]) != "" for r in body_rows)]

    str_rows = [[_fmt_native(vals[r][c]) for c in keep] for r in body_rows]
    h_rows = [[_fmt_native(hfill[r][c]) for c in keep] for r in body_rows]
    val_rows = [[vals[r][c] for c in keep] for r in body_rows]

    header, body = split_header(h_rows)
    n_head = len(str_rows) - len(body)
    body_str, body_val = str_rows[n_head:], val_rows[n_head:]

    text_cells = [s for s, v in zip(sum(body_str, []), sum(body_val, [])) if s and isinstance(v, str)]
    text_style = detect_number_style(text_cells) if text_cells else "vn"
    values = [[float(v) if _is_native_number(v) else parse_number(s, text_style)
               for s, v in zip(rs, rv)] for rs, rv in zip(body_str, body_val)]

    title = ""
    lead_text = [as_text(r) for r in lead]
    if lead_text and len(lead_text[-1]) <= 160:
        title = lead_text.pop()
    table = {
        "title": title or "Bảng (không có tiêu đề)", "bbox": None, "source": "excel",
        "header": header, "rows": body_str, "values": values,
        "number_style": "native" if any(_is_native_number(v) for r in body_val for v in r) else text_style,
    }
    return lead_text, table, [as_text(r) for r in trail], title


def read_excel(path: Path) -> List[Dict[str, Any]]:
    import openpyxl
    wb_v = openpyxl.load_workbook(str(path), data_only=True)
    wb_f = openpyxl.load_workbook(str(path), data_only=False)
    out = []
    for idx, ws_v in enumerate(wb_v.worksheets, 1):
        ws_f = wb_f[ws_v.title]
        vals, hfill, uncalc = _sheet_matrix(ws_v, ws_f)
        nonempty = [r for r, row in enumerate(vals) if any(_fmt_native(v) != "" for v in row)]
        hidden_state = ws_v.sheet_state != "visible"
        name = ws_v.title + (" (sheet ẩn)" if hidden_state else "")
        if not nonempty:
            out.append({"markdown": "", "tables": [], "metadata": {"ten_ban_ve": name}, "low_confidence": [],
                        "layout": "office_excel", "warnings": [f"Sheet '{ws_v.title}' trống."], "empty": True})
            continue
        parts: List[str] = []
        tables: List[Dict[str, Any]] = []
        k = 0
        pending = ""
        for block in _blocks(nonempty):
            lead, table, trail, title = _excel_block(vals, hfill, block)
            if table is None and len(lead) == 1 and len(lead[0]) <= 160 and not trail:
                pending = lead[0]           # dòng tiêu đề đứng riêng, cách bảng một dòng trống
                continue
            if pending:
                if table and not title:
                    title = pending
                    table["title"] = pending
                else:
                    parts.append(pending)
                pending = ""
            parts += [t for t in lead if t]
            if table:
                k += 1
                if not title:
                    table["title"] = f"Bảng {k}"
                parts.append(f"### {table['title']}\n\n{table_to_markdown(table)}")
                tables.append(table)
            parts += [t for t in trail if t]
        if pending:
            parts.append(pending)
        warnings = []
        if uncalc:
            warnings.append(f"{uncalc} ô công thức chưa có giá trị đã tính (file chưa được Excel mở và lưu) — "
                            "các ô này bị TRỐNG trong kết quả; mở file bằng Excel, lưu lại rồi chuyển đổi lại.")
        n_hidden_rows = sum(1 for d in ws_v.row_dimensions.values() if d.hidden)
        n_hidden_cols = sum(1 for d in ws_v.column_dimensions.values() if d.hidden)
        if n_hidden_rows or n_hidden_cols:
            warnings.append(f"Có dòng/cột bị ẩn ({n_hidden_rows} dòng, {n_hidden_cols} nhóm cột); "
                            "nội dung ẩn vẫn được đọc và đưa vào kết quả.")
        n_obj = len(getattr(ws_v, "_charts", [])) + len(getattr(ws_v, "_images", []))
        if n_obj:
            warnings.append(f"Sheet có {n_obj} biểu đồ/hình ảnh — chỉ đọc chữ và số trong ô, không đọc hình.")
        out.append({"markdown": "\n\n".join(parts), "tables": tables,
                    "metadata": {"ten_ban_ve": name}, "low_confidence": [],
                    "layout": "office_excel", "warnings": warnings})
    return out


# ─────────────────────────────────────────────────────────────────────────────
# WORD
# ─────────────────────────────────────────────────────────────────────────────
def _iter_body(parent_el, doc):
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in parent_el.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)
        elif child.tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from _iter_body(content, doc)


def _word_table_rows(table) -> List[List[str]]:
    rows = []
    try:
        for row in table.rows:
            prev_tc, cells = None, []
            for cell in row.cells:
                same = prev_tc is not None and cell._tc is prev_tc   # gộp ngang: chỉ giữ ô đầu
                cells.append("" if same else _ws(cell.text))
                prev_tc = cell._tc
            rows.append(cells)
    except Exception:
        from docx.oxml.ns import qn
        rows = []
        for tr in table._tbl.iter(qn("w:tr")):
            rows.append([_ws("".join(t.text or "" for t in tc.iter(qn("w:t")))) for tc in tr.iter(qn("w:tc"))])
    return rows


def read_word(path: Path) -> List[Dict[str, Any]]:
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(str(path))
    parts: List[str] = []
    tables: List[Dict[str, Any]] = []
    k = 0
    for item in _iter_body(doc.element.body, doc):
        if item.__class__.__name__ == "Paragraph":
            text = _ws(item.text)
            if not text:
                continue
            style = ""
            try:
                style = (item.style.name or "").lower()
            except Exception:
                pass
            ppr = item._p.pPr
            numpr = ppr.numPr if ppr is not None else None
            m = re.match(r"(?:heading|tiêu đề)\s*(\d)", style)
            if style == "title" or (m and int(m.group(1)) <= 2):
                parts.append(f"### {text}")
            elif m:
                parts.append(f"#### {text}")
            elif numpr is not None or "list" in style:
                lvl = int(numpr.ilvl.val) if numpr is not None and numpr.ilvl is not None else 0
                parts.append("  " * lvl + f"- {text}")
            else:
                parts.append(text)
        else:
            rows = _word_table_rows(item)
            title = ""
            if parts and not parts[-1].startswith(("#", "-", "  ")) and len(parts[-1]) <= 160 \
                    and fold_key(parts[-1]).startswith("BANG"):
                title = parts.pop()
            t = make_table(rows, title, None, "word")
            if t is None:
                lines = [" | ".join(c for c in r if c) for r in rows if any(r)]
                parts.append("\n".join(lines))
                continue
            k += 1
            if not title:
                t["title"] = f"Bảng {k}"
            parts.append(f"### {t['title']}\n\n{table_to_markdown(t)}")
            tables.append(t)

    boxes, seen = [], set()
    for tb in doc.element.body.iter(qn("w:txbxContent")):
        text = _ws(" ".join((t.text or "") for t in tb.iter(qn("w:t"))))
        if text and text not in seen:          # AlternateContent lặp mỗi hộp hai lần
            seen.add(text)
            boxes.append(text)
    if boxes:
        parts.append("### CHỮ TRONG HỘP VĂN BẢN\n\n" + "\n\n".join(f"- {b}" for b in boxes))

    warnings = []
    try:
        n_img = len(doc.inline_shapes)
    except Exception:
        n_img = 0
    if n_img:
        warnings.append(f"Có {n_img} hình ảnh trong file — chỉ đọc chữ và bảng, không đọc hình "
                        "(hình chụp/scan bản vẽ cần đưa ra file PDF hoặc ảnh để OCR).")
    if any(True for _ in doc.element.body.iter(qn("w:footnoteReference"))):
        warnings.append("File có chú thích cuối trang (footnote) chưa được đọc.")
    return [{"markdown": "\n\n".join(parts), "tables": tables, "metadata": {}, "low_confidence": [],
             "layout": "office_word", "warnings": warnings}]


def read_office(path) -> List[Dict[str, Any]]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext in LEGACY_EXT:
        raise _legacy_error(ext)
    if ext in WORD_EXT:
        return read_word(path)
    if ext in EXCEL_EXT:
        return read_excel(path)
    raise ValueError(f"Định dạng {ext} chưa được hỗ trợ.")
